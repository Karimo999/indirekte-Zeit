#!/usr/bin/env python3
"""Erzeugt Test-Zeitnachweise (Excel, Aufbau wie der Kilanka-Export) mit erfundenen Namen + erwartete Werte."""
import datetime, json, os
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ZIEL = os.path.join(HERE, '..', '..', 'test', 'faelle')
WT = ['Mo.', 'Di.', 'Mi.', 'Do.', 'Fr.', 'Sa.', 'So.']


def rhu(a, b):
    return (2 * a + b) // (2 * b)


def centi(minuten):
    return rhu(minuten * 100, 60)


def de(c):
    return f"{c // 100},{c % 100:02d}"


def minuten(von, bis):
    h1, m1 = map(int, von.split(':')); h2, m2 = map(int, bis.split(':'))
    return (h2 * 60 + m2) - (h1 * 60 + m1)


def erwartung(eintraege, von, bis):
    """eintraege: (datum, minuten, klient) im Zeitraum"""
    klienten = {}
    for d, m, k in eintraege:
        if k and m > 0 and von <= d <= bis:
            klienten.setdefault(k, []).append((d, m))
    tage_monat = (bis - von).days + 1
    voll = von.day == 1 and (bis + datetime.timedelta(days=1)).day == 1
    if voll:
        monate = (bis.year - von.year) * 12 + bis.month - von.month + 1
        grenze = 1302 * monate
    else:
        grenze = rhu(300 * tage_monat, 7)
        if tage_monat <= 31 and grenze > 1302:
            grenze = 1302
    seg = pin = 0
    for k, liste in klienten.items():
        seg += len({d for d, _ in liste}) * 20
        direkt = centi(sum(m for _, m in liste))
        pin += min(rhu(direkt, 5), grenze)
    return {'segeberg': de(seg), 'pinneberg': de(pin)}


def fall1():
    """Kilanka-Format: Datum als Text, Stunden dezimal, mehrere Eintraege am selben Tag"""
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Zeitnachweis'
    ws['A1'] = 'Mitarbeiter:'; ws['C1'] = 'Muster, Mia'
    ws['A2'] = 'Von:'; ws['C2'] = '01.09.2026'
    ws['A3'] = 'Bis:'; ws['C3'] = '30.09.2026'
    for i, h in enumerate(['Datum', 'Arbeitszeit', 'Nachtanteil', 'Pause', 'Stunden', 'Klient', 'Kommentar']):
        ws.cell(row=5, column=i + 1, value=h)
    daten = [
        (1, '08:00', '12:30', 'Kind, Anton', ''), (2, '08:00', '13:15', 'Kind, Anton', ''),
        (3, None, 4.494, '', 'Krank'), (4, '08:00', '11:00', 'Kind, Berta', ''),
        (7, '08:00', '12:00', 'Kind, Anton', ''), (7, '13:00', '15:30', 'Kind, Anton', 'Nachmittag'),
        (8, '08:00', '12:00', 'Kind, Anton', ''), (8, '13:00', '16:00', 'Kind, Berta', ''),
        (9, None, 5.35, '', 'SB Termin'), (10, '08:10', '13:05', 'Kind, Berta', ''),
        (11, '15:00', '17:00', '', 'Freizeitassistenz'), (14, '08:00', '14:00', 'Kind, Anton', ''),
    ]
    eintraege = []
    r = 6
    for tag, von, bis, klient, kommentar in daten:
        d = datetime.date(2026, 9, tag)
        ws.cell(row=r, column=1, value=f"{d:%d.%m.%Y} {WT[d.weekday()]}")
        if von:
            m = minuten(von, bis)
            ws.cell(row=r, column=2, value=f"{von}-{bis}")
            ws.cell(row=r, column=5, value=m / 60)
        else:
            m = round(bis * 60)
            ws.cell(row=r, column=5, value=bis)
        ws.cell(row=r, column=6, value=klient)
        ws.cell(row=r, column=7, value=kommentar)
        eintraege.append((d, m, klient))
        r += 1
    ws.cell(row=r, column=5, value=f'=SUM(E6:E{r - 1})')
    ws.cell(row=r + 2, column=1, value='Übersicht')
    ws.cell(row=r + 3, column=1, value='Ist-Stunden'); ws.cell(row=r + 3, column=2, value=12.5)
    ws.cell(row=r + 6, column=1, value='Unterschrift')
    ws.cell(row=r + 9, column=1, value='29.10.2026, Muster, Mia')
    name = 'zeitnachweis_muster_mia_2026-09.xlsx'
    wb.save(os.path.join(ZIEL, name))
    e = erwartung(eintraege, datetime.date(2026, 9, 1), datetime.date(2026, 9, 30))
    return name, {'mitarbeiter': 'Muster, Mia', 'zeitraum': '01.09.2026 - 30.09.2026', 'klienten': 2, **e}


def fall2():
    """Datum als Excel-Datum, Stunden als h:mm-Text, Teilzeitraum, Eintrag ausserhalb des Zeitraums"""
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Tabelle1'
    ws['A1'] = 'Mitarbeiterin'; ws['B1'] = 'Beispiel, Ben'
    ws['A2'] = 'Von'; ws['B2'] = datetime.date(2026, 9, 15); ws['B2'].number_format = 'DD.MM.YYYY'
    ws['A3'] = 'Bis'; ws['B3'] = datetime.date(2026, 9, 30); ws['B3'].number_format = 'DD.MM.YYYY'
    for i, h in enumerate(['Datum', 'Stunden', 'Klient/in', 'Bemerkung']):
        ws.cell(row=5, column=i + 1, value=h)
    eintraege = []
    r = 6
    for tag, hm in [(14, '5:00'), (15, '6:00'), (16, '6:00'), (17, '6:00'), (18, '6:00'), (21, '6:00'), (22, '6:00')]:
        d = datetime.date(2026, 9, tag)
        c = ws.cell(row=r, column=1, value=d); c.number_format = 'DD.MM.YYYY'
        ws.cell(row=r, column=2, value=hm)
        ws.cell(row=r, column=3, value='Kind, Carla')
        h, m = map(int, hm.split(':'))
        eintraege.append((d, h * 60 + m, 'Kind, Carla'))
        r += 1
    name = 'zeitnachweis_beispiel_ben_teilmonat.xlsx'
    wb.save(os.path.join(ZIEL, name))
    e = erwartung(eintraege, datetime.date(2026, 9, 15), datetime.date(2026, 9, 30))
    return name, {'mitarbeiter': 'Beispiel, Ben', 'zeitraum': '15.09.2026 - 30.09.2026', 'klienten': 1, **e}


def fall3():
    """gleicher Mitarbeiter und Monat wie ein Leistungsnachweis (Vorrang des Zeitnachweises in der Ansicht)"""
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Zeitnachweis'
    ws['A1'] = 'Mitarbeiter:'; ws['C1'] = 'Schulbegleiter, Max'
    ws['A2'] = 'Von:'; ws['C2'] = '01.09.2026'
    ws['A3'] = 'Bis:'; ws['C3'] = '30.09.2026'
    for i, h in enumerate(['Datum', 'Arbeitszeit', 'Nachtanteil', 'Pause', 'Stunden', 'Klient', 'Kommentar']):
        ws.cell(row=5, column=i + 1, value=h)
    eintraege = []
    for r, tag in enumerate([1, 2, 3], start=6):
        d = datetime.date(2026, 9, tag)
        ws.cell(row=r, column=1, value=f"{d:%d.%m.%Y} {WT[d.weekday()]}")
        ws.cell(row=r, column=2, value='08:00-13:00')
        ws.cell(row=r, column=5, value=5.0)
        ws.cell(row=r, column=6, value='Begleitkind, Maxime')
        eintraege.append((d, 300, 'Begleitkind, Maxime'))
    name = 'zeitnachweis_schulbegleiter_max_2026-09.xlsx'
    wb.save(os.path.join(ZIEL, name))
    e = erwartung(eintraege, datetime.date(2026, 9, 1), datetime.date(2026, 9, 30))
    return name, {'mitarbeiter': 'Schulbegleiter, Max', 'zeitraum': '01.09.2026 - 30.09.2026', 'klienten': 1, **e}


if __name__ == '__main__':
    erwartet = dict([fall1(), fall2(), fall3()])
    with open(os.path.join(ZIEL, 'erwartet_zeitnachweis.json'), 'w') as f:
        json.dump(erwartet, f, ensure_ascii=False, indent=1)
    print(json.dumps(erwartet, ensure_ascii=False, indent=1))
